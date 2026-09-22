import io
import time
import requests
from pathlib import Path
from django.core.files.storage import FileSystemStorage
from docx import Document as DocxDocument
from django.conf import settings
from django.contrib import messages
from django.contrib.auth import authenticate, login, logout, get_user_model
from django.contrib.auth.decorators import login_required
from django.core.paginator import Paginator
from django.core.mail import send_mail
from django.template.loader import render_to_string
from django.urls import reverse
from django.contrib.auth.tokens import default_token_generator
from django.utils.encoding import force_bytes
from django.utils.http import urlsafe_base64_encode
from django.db.models import Q
from django.http import HttpResponse
from django.shortcuts import render, redirect, get_object_or_404
from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from .forms import (
    LoginForm,
    SignUpForm,
    AdminUserCreateForm,
    AdminUserUpdateForm,
    PasswordRecoveryForm,
)
from .models import Analysis, StudyDocument


def login_view(request):
    if request.user.is_authenticated:
        return redirect("detector:dashboard")
    form = LoginForm(request, data=request.POST or None)
    if request.method == "POST" and form.is_valid():
        login(request, form.get_user())
        return redirect("detector:dashboard")
    return render(request, "login.html", {"form": form})



def password_reset_request(request):
    if request.user.is_authenticated:
        return redirect("detector:dashboard")

    form = PasswordRecoveryForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        user = form.user
        uid = urlsafe_base64_encode(force_bytes(user.pk))
        token = default_token_generator.make_token(user)
        reset_path = reverse(
            "detector:password_reset_confirm",
            kwargs={"uidb64": uid, "token": token},
        )
        base_url = getattr(settings, "PUBLIC_APP_URL", "").rstrip("/")
        if base_url:
            reset_url = f"{base_url}{reset_path}"
        else:
            reset_url = f"{'https' if request.is_secure() else 'http'}://{request.get_host()}{reset_path}"
        context = {
            "email": user.email,
            "user": user,
            "domain": request.get_host(),
            "site_name": "TDRDOC-SCAN",
            "protocol": "https" if request.is_secure() else "http",
            "uid": uid,
            "token": token,
            "reset_url": reset_url,
        }
        subject = render_to_string("password_reset_subject.txt", context).strip()
        message = render_to_string("password_reset_email.txt", context)

        required_email_settings = (
            getattr(settings, "EMAIL_HOST", ""),
            getattr(settings, "EMAIL_HOST_USER", ""),
            getattr(settings, "EMAIL_HOST_PASSWORD", ""),
            getattr(settings, "DEFAULT_FROM_EMAIL", ""),
        )
        if not all(required_email_settings):
            form.add_error(
                None,
                "L'envoi e-mail n'est pas configuré. Renseignez EMAIL_HOST, EMAIL_HOST_USER, EMAIL_HOST_PASSWORD et DEFAULT_FROM_EMAIL dans le fichier .env.",
            )
        else:
            try:
                send_mail(
                    subject,
                    message,
                    settings.DEFAULT_FROM_EMAIL,
                    [user.email],
                    fail_silently=False,
                )
            except Exception as exc:
                if settings.DEBUG:
                    form.add_error(None, f"Échec de l'envoi SMTP : {exc}")
                else:
                    form.add_error(
                        None,
                        "Le message de récupération n'a pas pu être envoyé. Vérifiez la configuration e-mail du serveur.",
                    )
            else:
                return redirect("detector:password_reset_done")

    return render(request, "password_reset.html", {"form": form})


def password_reset_done(request):
    return render(request, "password_reset_done.html")


def password_reset_confirm(request, uidb64, token):
    from django.contrib.auth.views import PasswordResetConfirmView
    view = PasswordResetConfirmView.as_view(
        template_name="password_reset_confirm.html",
        success_url="/mot-de-passe-reinitialise/",
    )
    return view(request, uidb64=uidb64, token=token)


def password_reset_complete(request):
    return render(request, "password_reset_complete.html")



def signup_view(request):
    if request.user.is_authenticated:
        return redirect("detector:dashboard")
    form = SignUpForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        user = form.save()
        messages.success(request, "Votre compte a été créé. Vous pouvez maintenant vous connecter.")
        return redirect("detector:login")
    return render(request, "signup.html", {"form": form})


def logout_view(request):
    logout(request)
    return redirect("detector:login")


def _get_uploaded_docx(request):
    candidates = ("document", "tdr_document", "report_document", "file")
    for field_name in candidates:
        uploaded = request.FILES.get(field_name)
        if uploaded and uploaded.name:
            return uploaded
    for uploaded in request.FILES.values():
        if uploaded and uploaded.name:
            return uploaded
    return None


def _validate_docx(uploaded):
    if not uploaded or not uploaded.name:
        return False
    return uploaded.name.strip().lower().endswith(".docx")


@login_required
def dashboard(request):
    base = Analysis.objects.filter(user=request.user).order_by("-created_at")
    q = request.GET.get("q", "").strip()
    decision = request.GET.get("decision", "all")
    mode = request.GET.get("mode", "all")
    if q:
        base = base.filter(Q(document_name__icontains=q) | Q(result_json__best_source__icontains=q))
    if decision in ("SIMILAIRE", "DIFFERENT", "A EXAMINER"):
        base = base.filter(decision=decision)
    if mode in ("duplicate", "plagiarism"):
        base = base.filter(mode=mode)
    paginator = Paginator(base, 10)
    page_obj = paginator.get_page(request.GET.get("page", 1))

    total = Analysis.objects.filter(user=request.user).count()
    similar = Analysis.objects.filter(user=request.user, decision="SIMILAIRE").count()
    different = Analysis.objects.filter(user=request.user, decision="DIFFERENT").count()
    to_review = Analysis.objects.filter(user=request.user, decision="A EXAMINER").count()
    tdr_checks = Analysis.objects.filter(user=request.user, mode="duplicate").count()
    report_checks = Analysis.objects.filter(user=request.user, mode="plagiarism").count()
    registry_total = StudyDocument.objects.count()
    registry_tdr = StudyDocument.objects.filter(document_type="TDR").count()
    registry_rapport = StudyDocument.objects.filter(document_type="RAPPORT").count()
    has_admin = get_user_model().objects.filter(is_staff=True, is_active=True).exists()
    return render(request, "dashboard.html", {
        "recent": page_obj.object_list, "page_obj": page_obj, "total_filtered": paginator.count,
        "q": q, "selected_decision": decision, "selected_mode": mode,
        "total": total, "similar": similar, "different": different, "to_review": to_review,
        "tdr_checks": tdr_checks, "report_checks": report_checks,
        "registry_total": registry_total, "registry_tdr": registry_tdr, "registry_rapport": registry_rapport,
        "has_admin": has_admin,
    })



def _save_uploaded_document(uploaded):
    """Enregistre le DOCX reçu afin de pouvoir le consulter depuis l'historique."""
    media_root = Path(settings.BASE_DIR) / "media" / "documents"
    media_root.mkdir(parents=True, exist_ok=True)
    storage = FileSystemStorage(location=str(media_root))
    uploaded.seek(0)
    return str(media_root / storage.save(uploaded.name, uploaded))


def _read_saved_document(path):
    """Lit un document DOCX précédemment enregistré et retourne son contenu."""
    if not path:
        return ""
    file_path = Path(path)
    if not file_path.exists() or file_path.suffix.lower() != ".docx":
        return ""
    doc = DocxDocument(file_path)
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\\n\\n".join(parts)

def _read_document_from_corpus(filename, mode):
    """Recherche le document analysé dans le corpus et retourne son texte complet."""
    root = Path(settings.BASE_DIR).resolve().parent
    folder = root / "donnees" / ("TDR" if mode == "duplicate" else "Rapport d'etude")
    if not folder.exists():
        return ""
    target = Path(filename).name.strip().lower()
    for path in folder.rglob("*.docx"):
        if path.name.strip().lower() == target:
            doc = DocxDocument(path)
            parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
            return "\\n\\n".join(parts)
    return ""


@login_required
def analysis_detail(request, pk):
    analysis = get_object_or_404(Analysis, pk=pk, user=request.user)
    result = analysis.result_json or {}
    document_text = _read_saved_document((analysis.result_json or {}).get("uploaded_path")) or _read_document_from_corpus(analysis.document_name, analysis.mode)
    return render(request, "analysis_detail.html", {
        "analysis": analysis,
        "result": result,
        "document_text": document_text,
    })


@login_required
def reception(request):
    if request.method == "GET":
        return render(request, "reception.html")
    mode = request.POST.get("mode", "tdr")
    uploaded = _get_uploaded_docx(request)
    if not _validate_docx(uploaded):
        messages.error(request, "Veuillez sélectionner un document DOCX valide.")
        return redirect("detector:reception")
    endpoint = "/api/detect/duplicate" if mode == "tdr" else "/api/detect/plagiarism"
    start = time.perf_counter()
    try:
        response = requests.post(
            f"{settings.FASTAPI_URL}{endpoint}",
            files={"file": (uploaded.name, uploaded.file, uploaded.content_type or "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            timeout=600,
        )
        response.raise_for_status()
        result = response.json()
        duration = round((time.perf_counter() - start) * 1000)
        result["duration_ms"] = duration
        analysis = Analysis.objects.create(
            user=request.user, document_name=uploaded.name,
            mode="duplicate" if mode == "tdr" else "plagiarism",
            decision=result.get("decision", ""), hybrid_score=result.get("hybrid_score"),
            semantic_score=result.get("semantic_score"), tfidf_score=result.get("tfidf_score"),
            novelty_score=result.get("novelty_score"), result_json=result, duration_ms=duration,
        )
        request.session["last_result"] = result
        request.session["last_analysis_id"] = analysis.id
        request.session["last_document_name"] = uploaded.name
        if mode == "tdr" and result.get("decision") == "DIFFERENT":
            request.session["validated_tdr_name"] = uploaded.name
            messages.success(request, "TDR accepté : aucune similarité suffisante n'a été trouvée. Le rapport associé peut maintenant être réceptionné.")
            return redirect("detector:reception")
        if mode == "tdr" and result.get("decision") == "SIMILAIRE":
            messages.error(request, "TDR non retenu : un document similaire existe déjà dans la base.")
        else:
            messages.info(request, "Rapport analysé. Consultez le détail des résultats.")
        return redirect("detector:resultats")
    except requests.RequestException as exc:
        messages.error(request, f"Le moteur d'analyse est indisponible : {exc}")
    return redirect("detector:reception")



@login_required
def upload_report_for_tdr(request, pk):
    tdr = get_object_or_404(Analysis, pk=pk, user=request.user, mode="duplicate", decision="DIFFERENT")
    if request.method == "GET":
        return render(request, "report_upload.html", {"tdr": tdr})
    uploaded = _get_uploaded_docx(request)
    if not _validate_docx(uploaded):
        messages.error(request, "Veuillez sélectionner un rapport DOCX valide.")
        return redirect("detector:upload_report_for_tdr", pk=tdr.pk)
    start = time.perf_counter()
    try:
        response = requests.post(
            f"{settings.FASTAPI_URL}/api/detect/plagiarism",
            files={"file": (uploaded.name, uploaded.file, uploaded.content_type or "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            timeout=600,
        )
        response.raise_for_status()
        result = response.json()
        result["uploaded_path"] = _save_uploaded_document(uploaded)
        result["linked_tdr_analysis_id"] = tdr.id
        result["linked_tdr_name"] = tdr.document_name
        duration = round((time.perf_counter() - start) * 1000)
        result["duration_ms"] = duration
        analysis = Analysis.objects.create(
            user=request.user, document_name=uploaded.name, mode="plagiarism",
            decision=result.get("decision", ""), hybrid_score=result.get("hybrid_score"),
            semantic_score=result.get("semantic_score"), tfidf_score=result.get("tfidf_score"),
            novelty_score=result.get("novelty_score"), result_json=result, duration_ms=duration,
        )
        request.session["last_result"] = result
        request.session["last_analysis_id"] = analysis.id
        request.session["last_document_name"] = uploaded.name
        messages.success(request, "Rapport associé au TDR et analysé avec succès.")
        return redirect("detector:resultats")
    except requests.RequestException as exc:
        messages.error(request, f"Le moteur d'analyse est indisponible : {exc}")
        return redirect("detector:upload_report_for_tdr", pk=tdr.pk)

@login_required
def analyser(request):
    if request.method != "POST":
        return render(request, "analyser.html")
    uploaded = _get_uploaded_docx(request)
    mode = request.POST.get("mode", "plagiarism")
    if mode not in ("plagiarism", "duplicate"):
        messages.error(request, "Mode d'analyse invalide.")
        return redirect("detector:analyser")
    if not _validate_docx(uploaded):
        messages.error(request, "Veuillez sélectionner un document DOCX valide.")
        return redirect("detector:analyser")
    endpoint = "/api/detect/duplicate" if mode == "duplicate" else "/api/detect/plagiarism"
    start = time.perf_counter()
    try:
        response = requests.post(
            f"{settings.FASTAPI_URL}{endpoint}",
            files={"file": (uploaded.name, uploaded.file, uploaded.content_type or "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            timeout=600,
        )
        response.raise_for_status()
        result = response.json()
        duration = round((time.perf_counter() - start) * 1000)
        result["duration_ms"] = duration
        analysis = Analysis.objects.create(
            user=request.user, document_name=uploaded.name, mode=mode,
            decision=result.get("decision", ""), hybrid_score=result.get("hybrid_score"),
            semantic_score=result.get("semantic_score"), tfidf_score=result.get("tfidf_score"),
            novelty_score=result.get("novelty_score"), result_json=result, duration_ms=duration,
        )
        request.session["last_result"] = result
        request.session["last_analysis_id"] = analysis.id
        request.session["last_document_name"] = uploaded.name
        return redirect("detector:resultats")
    except requests.RequestException as exc:
        messages.error(request, f"Le moteur FastAPI est indisponible : {exc}")
    except Exception as exc:
        messages.error(request, f"Erreur pendant l'analyse : {exc}")
    return render(request, "analyser.html")


@login_required
def resultats(request):
    """Historique complet des analyses de l'utilisateur, avec filtres et pagination."""
    analyses = Analysis.objects.filter(user=request.user).order_by("-created_at")
    q = request.GET.get("q", "").strip()
    decision = request.GET.get("decision", "all")
    mode = request.GET.get("mode", "all")
    if q:
        analyses = analyses.filter(Q(document_name__icontains=q) | Q(result_json__best_source__icontains=q))
    if decision in ("SIMILAIRE", "DIFFERENT", "A EXAMINER"):
        analyses = analyses.filter(decision=decision)
    if mode in ("duplicate", "plagiarism"):
        analyses = analyses.filter(mode=mode)
    paginator = Paginator(analyses, 10)
    page_obj = paginator.get_page(request.GET.get("page", 1))
    return render(request, "resultats.html", {
        "analyses": page_obj.object_list,
        "page_obj": page_obj,
        "total_filtered": paginator.count,
        "q": q,
        "selected_decision": decision,
        "selected_mode": mode,
    })


@login_required
def export_excel(request):
    analysis = Analysis.objects.filter(user=request.user).order_by("-created_at").first()
    if not analysis:
        return HttpResponse("Aucune analyse à exporter.", status=404)
    result = analysis.result_json or {}
    wb = Workbook(); ws = wb.active; ws.title = "Analyse"
    for row in [["Détecteur documentaire", "DOCSEC"], ["Document", analysis.document_name], ["Mode", analysis.get_mode_display()], ["Décision", analysis.decision], ["Score TF-IDF", analysis.tfidf_score], ["Score sémantique", analysis.semantic_score], ["Score hybride", analysis.hybrid_score], ["Score de nouveauté", analysis.novelty_score], ["Durée (ms)", analysis.duration_ms], ["Meilleure source", result.get("best_source", "")]]: ws.append(row)
    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"); response["Content-Disposition"] = 'attachment; filename="rapport_analyse.xlsx"'; wb.save(response); return response


@login_required
def export_pdf(request):
    analysis = Analysis.objects.filter(user=request.user).order_by("-created_at").first()
    if not analysis: return HttpResponse("Aucune analyse à exporter.", status=404)
    result = analysis.result_json or {}; buffer = io.BytesIO(); doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36); styles = getSampleStyleSheet()
    story = [Paragraph("Rapport d'analyse documentaire — DOCSEC", styles["Title"]), Spacer(1, 12)]
    data = [["Document", analysis.document_name], ["Type", analysis.get_mode_display()], ["Décision", analysis.decision], ["Score TF-IDF", str(analysis.tfidf_score)], ["Score sémantique", str(analysis.semantic_score)], ["Score hybride", str(analysis.hybrid_score)], ["Nouveauté", str(analysis.novelty_score)], ["Meilleure source", str(result.get("best_source", ""))]]
    table = Table(data, colWidths=[150, 350]); table.setStyle(TableStyle([("BACKGROUND", (0,0), (0,-1), colors.HexColor("#eef2f7")), ("GRID", (0,0), (-1,-1), .5, colors.grey), ("VALIGN", (0,0), (-1,-1), "TOP"), ("PADDING", (0,0), (-1,-1), 7)])); story.append(table); story.append(Spacer(1, 18)); story.append(Paragraph("Sources proches", styles["Heading2"]))
    for source in result.get("sources", [])[:10]: story.extend([Paragraph(f"{source.get('source','')} — score {source.get('hybrid_score','')}", styles["BodyText"]), Spacer(1, 5)])
    doc.build(story); response = HttpResponse(buffer.getvalue(), content_type="application/pdf"); response["Content-Disposition"] = 'attachment; filename="rapport_analyse.pdf"'; return response





@login_required
def administration_recovery(request):
    """Récupération contrôlée de l'accès administrateur lorsqu'il n'existe plus aucun administrateur actif."""
    User = get_user_model()
    active_admin_exists = User.objects.filter(is_staff=True, is_active=True).exists()
    if active_admin_exists:
        return redirect("detector:administration")

    if not settings.DEBUG and not getattr(settings, "ALLOW_ADMIN_BOOTSTRAP", False):
        messages.error(
            request,
            "Aucun administrateur actif n'est configuré. Le mode de récupération est désactivé sur cette installation.",
        )
        return redirect("detector:dashboard")

    if request.method == "POST":
        user = User.objects.get(pk=request.user.pk)
        user.is_staff = True
        user.is_active = True
        user.save(update_fields=["is_staff", "is_active"])
        messages.success(
            request,
            "Votre compte a retrouvé le rôle Administrateur. Vous pouvez maintenant accéder à la console.",
        )
        return redirect("detector:administration")

    return render(request, "administration_recovery.html", {
        "username": request.user.get_username(),
    })

@login_required
def administration(request):
    User = get_user_model()
    if not request.user.is_staff:
        if not User.objects.filter(is_staff=True, is_active=True).exists() and (
            settings.DEBUG or getattr(settings, "ALLOW_ADMIN_BOOTSTRAP", False)
        ):
            return redirect("detector:administration_recovery")
        messages.error(request, "Accès réservé aux administrateurs.")
        return redirect("detector:dashboard")

    User = get_user_model()
    q = request.GET.get("q", "").strip()
    users = User.objects.order_by("-is_active", "username")
    if q:
        users = users.filter(
            Q(username__icontains=q)
            | Q(first_name__icontains=q)
            | Q(last_name__icontains=q)
            | Q(email__icontains=q)
        )

    if request.method == "POST":
        action = request.POST.get("action", "")

        if action == "create_user":
            form = AdminUserCreateForm(request.POST)
            if form.is_valid():
                user = form.save()
                messages.success(request, f"Le compte « {user.username} » a été créé.")
                return redirect("detector:administration")
            return render(request, "administration.html", {
                "create_user_form": form,
                "users": users,
                "stats": _admin_user_stats(User),
                "current_section": "users",
            })

        if action == "toggle_user":
            user = get_object_or_404(User, pk=request.POST.get("user_id"))
            if user.pk == request.user.pk:
                messages.error(request, "Vous ne pouvez pas désactiver votre propre compte depuis cette page.")
            else:
                user.is_active = not user.is_active
                user.save(update_fields=["is_active"])
                messages.success(request, f"Le compte « {user.username} » est maintenant {'actif' if user.is_active else 'inactif'}.")
            return redirect("detector:administration")

        if action == "delete_user":
            user = get_object_or_404(User, pk=request.POST.get("user_id"))
            if user.pk == request.user.pk:
                messages.error(request, "Vous ne pouvez pas supprimer votre propre compte.")
                return redirect("detector:administration")
            return redirect("detector:administration_utilisateur_supprimer", pk=user.pk)

    context = {
        "create_user_form": AdminUserCreateForm(),
        "users": users,
        "q": q,
        "stats": _admin_user_stats(User),
        "current_section": "users",
    }
    return render(request, "administration.html", context)


@login_required
def administration_utilisateur(request, pk):
    if not request.user.is_staff:
        messages.error(request, "Accès réservé aux administrateurs.")
        return redirect("detector:dashboard")

    User = get_user_model()
    target = get_object_or_404(User, pk=pk)
    profile_form = AdminUserUpdateForm(request.POST or None, instance=target)
    action = request.POST.get("action", "")
    if request.method == "POST":
        if action == "update_profile":
            profile_form = AdminUserUpdateForm(request.POST, instance=target)
            if profile_form.is_valid():
                if target.pk == request.user.pk and (
                    profile_form.cleaned_data["role"] != "admin"
                    or not profile_form.cleaned_data["is_active"]
                ):
                    profile_form.add_error(
                        "role",
                        "Votre propre compte doit rester Administrateur et actif.",
                    )
                else:
                    profile_form.save()
                    messages.success(request, f"Le compte « {target.username} » a été mis à jour.")
                    return redirect("detector:administration_utilisateur", pk=target.pk)

        elif action == "delete_user":
            if target.pk == request.user.pk:
                messages.error(request, "Vous ne pouvez pas supprimer votre propre compte.")
                return redirect("detector:administration_utilisateur", pk=target.pk)
            username = target.username
            return redirect("detector:administration_utilisateur_supprimer", pk=target.pk)

    return render(request, "administration_utilisateur.html", {
        "target_user": target,
        "profile_form": profile_form,
    })


@login_required
def administration_utilisateur_supprimer(request, pk):
    if not request.user.is_staff:
        messages.error(request, "Accès réservé aux administrateurs.")
        return redirect("detector:dashboard")

    User = get_user_model()
    target = get_object_or_404(User, pk=pk)
    if target.pk == request.user.pk:
        messages.error(request, "Vous ne pouvez pas supprimer votre propre compte.")
        return redirect("detector:administration")
    if target.is_staff and target.is_active and User.objects.filter(is_staff=True, is_active=True).count() <= 1:
        messages.error(request, "Impossible de supprimer le dernier administrateur actif.")
        return redirect("detector:administration")

    if request.method == "POST":
        username = target.username
        target.delete()
        messages.success(request, f"Le compte « {username} » a été supprimé.")
        return redirect("detector:administration")

    return render(request, "administration_utilisateur_supprimer.html", {
        "target_user": target,
    })


def _admin_user_stats(User):
    return {
        "users": User.objects.count(),
        "active": User.objects.filter(is_active=True).count(),
        "inactive": User.objects.filter(is_active=False).count(),
        "admins": User.objects.filter(is_staff=True).count(),
        "regular": User.objects.filter(is_staff=False).count(),
    }
