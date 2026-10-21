from pathlib import Path
from functools import lru_cache
import os

from django import template
from django.conf import settings
from docx import Document
from pypdf import PdfReader

from src.indexing.reference_index import load_reference_documents
from src.similarity.content_profile import build_content_profile

register = template.Library()


def _read_document(path: Path) -> str:
    if path.suffix.lower() == ".pdf":
        reader = PdfReader(str(path))
        return "\n\n".join((page.extract_text() or "").strip() for page in reader.pages if (page.extract_text() or "").strip())
    document = Document(path)
    parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


@lru_cache(maxsize=32)
def _profile_cached(analysis_id, uploaded_path, kind, threshold):
    path = (Path(settings.MEDIA_ROOT) / str(uploaded_path).replace("/", os.sep)).resolve()
    media_root = Path(settings.MEDIA_ROOT).resolve()
    try:
        path.relative_to(media_root)
    except ValueError:
        return None
    if not path.is_file():
        return None

    rows, embeddings = load_reference_documents(kind)
    if not rows:
        return None
    text = _read_document(path)
    return build_content_profile(text, rows, embeddings, lexical_weight=0.30, semantic_weight=0.70, threshold=float(threshold), kind=kind)


@register.simple_tag
def get_content_profile(analysis, result):
    """Calcule à l'affichage le profil V2 à partir du document et du corpus réel."""
    try:
        uploaded_path = (result or {}).get("uploaded_path", "")
        if not uploaded_path:
            return None
        kind = "duplicate" if getattr(analysis, "mode", "") == "duplicate" else "plagiarism"
        threshold = float((result or {}).get("threshold", 0.70 if kind == "duplicate" else 0.55))
        return _profile_cached(analysis.pk, uploaded_path, kind, threshold)
    except Exception as exc:
        return {"overall": {"existing_percent": 0.0, "new_percent": 0.0, "partial_percent": 0.0, "segments": 0}, "components": [], "interpretation": {"level": "NON DISPONIBLE", "recommendation": "Le profil de couverture n'a pas pu être calculé pour cette analyse.", "note": str(exc)}}
