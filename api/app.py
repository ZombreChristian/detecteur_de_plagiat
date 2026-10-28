
"""API FastAPI du moteur de détection de plagiat et de doublons."""

from contextlib import asynccontextmanager
from pathlib import Path
import html
import os
import re
import tempfile
import time

import numpy as np
from dotenv import load_dotenv
from docx import Document
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from pypdf import PdfReader
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
from starlette.concurrency import run_in_threadpool

ROOT = Path(__file__).resolve().parents[1]
load_dotenv(ROOT / "frontend" / ".env")

from src.indexing.reference_index import (
    ensure_reference_index,
    load_reference_documents,
)
from src.llm.ollama_rag import analyze_tdr_with_rag, OllamaRAGError
from src.passages.passage_engine import analyze_document_pair, adjusted_hybrid_score
from src.preprocessing.clean_text import clean_text
from src.similarity.contextual_engine import analyze_context
from src.similarity.model_manager import get_model, get_model_info


THRESHOLDS = {"plagiarism": 0.55}

LEXICAL_WEIGHT = float(os.getenv("SIMILARITY_LEXICAL_WEIGHT", "0.30"))
SEMANTIC_WEIGHT = float(os.getenv("SIMILARITY_SEMANTIC_WEIGHT", "0.70"))
TOP_SOURCES = int(os.getenv("SIMILARITY_REFERENCE_TOP_K", "10"))
TOP_PASSAGE_SOURCES = int(os.getenv("SIMILARITY_PASSAGE_TOP_K", "5"))

if abs((LEXICAL_WEIGHT + SEMANTIC_WEIGHT) - 1.0) > 1e-6:
    raise RuntimeError(
        "SIMILARITY_LEXICAL_WEIGHT + SIMILARITY_SEMANTIC_WEIGHT "
        "doit être égal à 1.0."
    )

if TOP_SOURCES < 1 or TOP_PASSAGE_SOURCES < 1:
    raise RuntimeError(
        "SIMILARITY_REFERENCE_TOP_K et SIMILARITY_PASSAGE_TOP_K "
        "doivent être supérieurs ou égaux à 1."
    )

_REFERENCE_CACHE = {}


def read_document(path: Path) -> str:
    """Extrait le texte d'un document PDF ou DOCX."""
    suffix = path.suffix.lower()

    try:
        if suffix == ".pdf":
            reader = PdfReader(str(path))
            pages = []

            for page in reader.pages:
                page_text = page.extract_text() or ""
                page_text = page_text.strip()

                if page_text:
                    pages.append(page_text)

            return "\n\n".join(pages)

        if suffix == ".docx":
            doc = Document(str(path))
            parts = [
                paragraph.text.strip()
                for paragraph in doc.paragraphs
                if paragraph.text.strip()
            ]

            for table in doc.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]

                    if any(cells):
                        parts.append(" | ".join(cells))

            return "\n\n".join(parts)

        raise HTTPException(
            status_code=400,
            detail="Format non pris en charge. Utilisez un fichier DOCX ou PDF.",
        )

    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail=f"Impossible de lire le document {path.name} : {exc}",
        ) from exc


def _build_reference_runtime_index(kind: str):
    """Construit et met en cache l'index TF-IDF des documents de référence."""
    if kind not in ("duplicate", "plagiarism"):
        raise HTTPException(
            status_code=404,
            detail="Mode inconnu. Utilisez plagiarism ou duplicate.",
        )

    cached = _REFERENCE_CACHE.get(kind)

    if cached is not None:
        return cached

    rows, embeddings = load_reference_documents(kind)

    if not rows:
        raise HTTPException(
            status_code=500,
            detail=f"Aucun document de référence disponible pour le mode {kind}.",
        )

    texts = [
        row.cleaned_text or row.extracted_text or ""
        for row in rows
    ]

    if not any(text.strip() for text in texts):
        raise HTTPException(
            status_code=500,
            detail=f"Les références du mode {kind} ne contiennent aucun texte exploitable.",
        )

    vectorizer = TfidfVectorizer(
        ngram_range=(1, 2),
        sublinear_tf=True,
    )

    try:
        matrix = vectorizer.fit_transform(texts)
    except ValueError as exc:
        raise HTTPException(
            status_code=500,
            detail=f"Impossible de construire l'index TF-IDF : {exc}",
        ) from exc

    embeddings = np.asarray(embeddings)

    if len(embeddings) != len(rows):
        raise HTTPException(
            status_code=500,
            detail=(
                "Incohérence dans l'index de référence : "
                "le nombre d'embeddings diffère du nombre de documents."
            ),
        )

    cached = {
        "rows": rows,
        "embeddings": embeddings,
        "vectorizer": vectorizer,
        "tfidf_matrix": matrix,
    }

    _REFERENCE_CACHE[kind] = cached
    return cached


def _warm_reference_runtime_indexes() -> None:
    """Préchauffe les index disponibles au démarrage."""
    for kind in ("duplicate", "plagiarism"):
        try:
            _build_reference_runtime_index(kind)
        except Exception as exc:
            print(f"[similarity] Index {kind} non disponible : {exc}")


def _semantic_scores(
    candidate_embedding: np.ndarray,
    reference_embeddings: np.ndarray,
) -> np.ndarray:
    """Calcule les similarités cosinus entre embeddings normalisés."""
    scores = np.matmul(reference_embeddings, candidate_embedding)
    return np.clip(scores, 0.0, 1.0)


def score_documents(candidate_text: str, kind: str):
    """Classe les documents de référence par score hybride."""
    if kind not in ("plagiarism", "duplicate"):
        raise HTTPException(
            status_code=404,
            detail="Mode inconnu. Utilisez plagiarism ou duplicate.",
        )

    candidate = clean_text(candidate_text)

    if not candidate:
        return []

    index = _build_reference_runtime_index(kind)

    vectorizer = index["vectorizer"]
    reference_tfidf = index["tfidf_matrix"]
    reference_embeddings = index["embeddings"]
    rows = index["rows"]

    candidate_tfidf = vectorizer.transform([candidate])

    lexical = cosine_similarity(
        candidate_tfidf,
        reference_tfidf,
    ).ravel()

    candidate_embedding = get_model().encode(
        [candidate],
        normalize_embeddings=True,
        show_progress_bar=False,
        convert_to_numpy=True,
    )[0]

    semantic = _semantic_scores(
        candidate_embedding,
        reference_embeddings,
    )

    hybrid = np.array(
        [
            adjusted_hybrid_score(
                candidate,
                rows[int(idx)].extracted_text
                or rows[int(idx)].cleaned_text,
                lexical[idx],
                semantic[idx],
                LEXICAL_WEIGHT,
                SEMANTIC_WEIGHT,
            )
            for idx in range(len(rows))
        ]
    )

    order = np.argsort(hybrid)[::-1]
    ranked = []

    for idx in order:
        row = rows[int(idx)]

        ranked.append(
            {
                "id": row.id,
                "source": Path(row.file_path).name or row.title,
                "title": row.title,
                "tfidf_score": round(float(lexical[idx]), 4),
                "semantic_score": round(float(semantic[idx]), 4),
                "hybrid_score": round(float(hybrid[idx]), 4),
                "_document": row,
            }
        )

    return ranked


def score_two_texts(candidate_text: str, source_text: str):
    """Compare directement deux textes."""
    candidate = clean_text(candidate_text)
    source = clean_text(source_text)

    if not candidate or not source:
        return None

    vectorizer = TfidfVectorizer(
        ngram_range=(1, 2),
        sublinear_tf=True,
    )

    try:
        matrix = vectorizer.fit_transform([candidate, source])
    except ValueError:
        return None

    document_lexical = float(
        cosine_similarity(matrix[0:1], matrix[1:2])[0, 0]
    )

    detail = analyze_document_pair(
        candidate,
        source,
        top_k=8,
        lexical_weight=LEXICAL_WEIGHT,
        semantic_weight=SEMANTIC_WEIGHT,
        document_lexical_score=document_lexical,
    )

    return {
        "tfidf_score": detail["tfidf_score"],
        "semantic_score": detail["semantic_score"],
        "hybrid_score": detail["hybrid_score"],
        "coverage": detail["coverage"],
        "passages_similaires": len(detail["matches"]),
    }


def decision(score, threshold):
    return "SIMILAIRE" if score >= threshold else "DIFFERENT"


def _contextual_explanation(contextual):
    """Produit une explication courte et lisible de l'identité de l'étude."""
    dimensions = contextual.get("dimensions", {})

    labels = {
        "objet": "objet",
        "secteur": "secteur",
        "activites": "activités",
        "zone": "zone géographique",
        "objectifs": "objectifs",
        "resultats": "résultats attendus",
    }

    similar = [
        labels[name]
        for name, data in dimensions.items()
        if name in labels and data.get("status") == "SIMILAIRE"
    ]

    different = [
        labels[name]
        for name, data in dimensions.items()
        if name in labels and data.get("status") == "DIFFERENT"
    ]

    if contextual.get("verdict") == "SIMILAIRE":
        if similar:
            return (
                "Les deux documents présentent une identité d'étude fortement "
                f"concordante sur {', '.join(similar)}. "
                "Les différences de formulation ou de structure ne suffisent pas "
                "à remettre en cause cette proximité."
            )

        return (
            "L'analyse contextuelle indique une forte proximité "
            "de l'identité de l'étude."
        )

    if contextual.get("verdict") == "DIFFERENT":
        if different:
            return (
                "L'analyse contextuelle relève une divergence importante sur "
                f"{', '.join(different)}. La similarité de structure ou de "
                "formulation ne suffit donc pas à considérer les deux études "
                "comme identiques."
            )

        return (
            "L'analyse contextuelle ne confirme pas une identité commune d'étude."
        )

    return (
        "Les documents présentent des éléments communs, mais l'identité de "
        "l'étude reste insuffisamment établie pour une décision automatique."
    )


def run_detection(
    candidate_text: str,
    kind: str,
    threshold: float | None = None,
):
    """Exécute le parcours de détection correspondant au type de document."""
    if kind not in ("duplicate", "plagiarism"):
        raise HTTPException(
            status_code=404,
            detail="Mode inconnu. Utilisez plagiarism ou duplicate.",
        )

    candidate = clean_text(candidate_text)

    if not candidate:
        raise HTTPException(
            status_code=400,
            detail="Le document ne contient aucun texte exploitable.",
        )

    ranked = score_documents(candidate, kind)

    if not ranked:
        if kind == "duplicate":
            raise HTTPException(
                status_code=409,
                detail=(
                    "Aucune référence TDR n'est indexée dans la base. "
                    "Impossible de comparer ce document."
                ),
            )

        raise HTTPException(
            status_code=500,
            detail=f"Aucun document de référence indexé pour le mode « {kind} ».",
        )

    best = ranked[0]

    # ---------------------------------------------------------
    # PARCOURS TDR : Qwen3 via Ollama est l'arbitre métier.
    # Les scores hybrides servent uniquement à récupérer
    # les références candidates pour l'analyse RAG.
    # ---------------------------------------------------------
    if kind == "duplicate":
        try:
            rag_analysis = analyze_tdr_with_rag(candidate, ranked)

        except OllamaRAGError as exc:
            raise HTTPException(
                status_code=503,
                detail={
                    "error": "TDR_RAG_UNAVAILABLE",
                    "message": str(exc),
                    "hint": (
                        "Démarrez Ollama et vérifiez que qwen3:8b est installé "
                        "(ollama list). Vérifiez aussi OLLAMA_MODEL "
                        "dans frontend/.env."
                    ),
                },
            ) from exc

        ranked_by_name = {
            item["source"]: item
            for item in ranked
        }

        used_names = rag_analysis.get("references_used", [])
        sources = []

        for source_name in used_names:
            item = ranked_by_name.get(source_name)

            if not item:
                continue

            sources.append(
                {
                    "source_id": item["id"],
                    "source": item["source"],
                    "title": item.get("title") or item["source"],
                    "retrieval_score": item["hybrid_score"],
                    "tfidf_score": item["tfidf_score"],
                    "semantic_score": item["semantic_score"],
                    "hybrid_score": item["hybrid_score"],
                    "matches": [],
                }
            )

        matched_reference = rag_analysis.get("matched_reference", "")
        matched_item = ranked_by_name.get(matched_reference)

        # Seule la référence explicitement retenue par Qwen est présentée
        # comme source de comparaison.
        if matched_item:
            sources = [
                source
                for source in sources
                if source["source"] == matched_reference
            ]
        else:
            sources = []

        final_decision = rag_analysis["decision"]

        return {
            "mode": "duplicate",
            "decision": final_decision,
            "threshold": None,
            "threshold_percent": None,
            "decision_provider": "ollama",
            "rag_analysis": rag_analysis,
            "lexical_weight": LEXICAL_WEIGHT,
            "semantic_weight": SEMANTIC_WEIGHT,
            "best_source": matched_reference if matched_item else "",
            "best_source_id": matched_item["id"] if matched_item else None,
            "candidate_document_html": html.escape(candidate),
            "tfidf_score": None,
            "semantic_score": None,
            "hybrid_score": None,
            "novelty_score": None,
            "contextual_analysis": None,
            "contextual_explanation": rag_analysis.get("explanation", ""),
            "sources": sources,
            "matches": [],
        }

    # ---------------------------------------------------------
    # PARCOURS RAPPORTS : analyse des passages et analyse
    # contextuelle. Aucun jugement IA supplémentaire n'est
    # effectué ici, car aucun moteur de jugement IA distinct
    # n'est défini dans ce fichier.
    # ---------------------------------------------------------
    if threshold is None:
        threshold = THRESHOLDS[kind]
    else:
        threshold = float(threshold)

        if threshold > 1.0:
            threshold /= 100.0

    if not 0.01 <= threshold <= 1.0:
        raise HTTPException(
            status_code=400,
            detail="Le seuil de décision doit être compris entre 1 % et 100 %.",
        )

    passage_groups = []

    for item in ranked[:TOP_PASSAGE_SOURCES]:
        document = item["_document"]
        source_text = document.extracted_text or document.cleaned_text

        detail = analyze_document_pair(
            candidate,
            source_text,
            top_k=8,
            lexical_weight=LEXICAL_WEIGHT,
            semantic_weight=SEMANTIC_WEIGHT,
            document_lexical_score=item["tfidf_score"],
        )

        source_html = html.escape(source_text)
        candidate_html = html.escape(candidate)

        unique_passages = []
        for match in detail["matches"]:
            passage = match["source_passage"]

            if passage and passage not in unique_passages:
                unique_passages.append(passage)

        candidate_passages = []
        for match in detail["matches"]:
            passage = match["candidate_passage"]

            if passage and passage not in candidate_passages:
                candidate_passages.append(passage)

        for passage in sorted(candidate_passages, key=len, reverse=True):
            escaped_passage = html.escape(passage)

            candidate_html = candidate_html.replace(
                escaped_passage,
                f'<mark class="docsec-evidence">{escaped_passage}</mark>',
            )

        for passage in sorted(unique_passages, key=len, reverse=True):
            escaped_passage = html.escape(passage)

            source_html = source_html.replace(
                escaped_passage,
                f'<mark class="docsec-evidence">{escaped_passage}</mark>',
            )

        contextual = analyze_context(candidate, source_text)

        passage_groups.append(
            {
                "source_id": item["id"],
                "source": item["source"],
                "document_score": detail["hybrid_score"],
                "tfidf_score": detail["tfidf_score"],
                "semantic_score": detail["semantic_score"],
                "hybrid_score": detail["hybrid_score"],
                "coverage": detail["coverage"],
                "matches": detail["matches"],
                "contextual_analysis": contextual,
                "contextual_verdict": contextual.get(
                    "verdict", "À EXAMINER"
                ),
                "contextual_confidence": contextual.get("confidence"),
                # Aucun jugement IA distinct n'est exécuté dans ce parcours.
                "ai_judgment": None,
                "ai_verdict": None,
                "ai_confidence": None,
                "passages_candidate": detail["passages_candidate"],
                "passages_source": detail["passages_source"],
                "source_document": source_text,
                "source_document_html": source_html,
                "candidate_document_html": candidate_html,
            }
        )

    evidence_groups = [
        group
        for group in passage_groups
        if group.get("matches")
    ]

    evidence_groups.sort(
        key=lambda group: (
            group.get("contextual_verdict") == "SIMILAIRE",
            group.get("hybrid_score", 0.0),
        ),
        reverse=True,
    )

    best_evidence = evidence_groups[0] if evidence_groups else None

    if best_evidence is not None:
        score_decision = decision(
            best_evidence["hybrid_score"],
            threshold,
        )

        contextual = best_evidence.get("contextual_analysis", {})
        context_verdict = best_evidence.get(
            "contextual_verdict",
            "À EXAMINER",
        )

        final_decision = (
            "DIFFERENT"
            if score_decision == "DIFFERENT"
            else context_verdict
        )

        result_tfidf = best_evidence["tfidf_score"]
        result_semantic = best_evidence["semantic_score"]
        result_hybrid = best_evidence["hybrid_score"]
        result_source = best_evidence["source"]
        result_source_id = best_evidence["source_id"]
        contextual_for_response = contextual
        explanation_for_response = _contextual_explanation(contextual)

    else:
        final_decision = "DIFFERENT"
        result_tfidf = 0.0
        result_semantic = 0.0
        result_hybrid = 0.0
        result_source = best["source"]
        result_source_id = best["id"]
        contextual_for_response = None

        explanation_for_response = (
            "Aucune correspondance de passages exploitable n'a été trouvée."
        )

    displayed_sources = sorted(
        passage_groups,
        key=lambda group: group.get("hybrid_score", 0.0),
        reverse=True,
    )

    return {
        "mode": kind,
        "decision": final_decision,
        "threshold": threshold,
        "threshold_percent": round(threshold * 100, 2),
        "decision_provider": "contextual_engine",
        "rag_analysis": None,
        "lexical_weight": LEXICAL_WEIGHT,
        "semantic_weight": SEMANTIC_WEIGHT,
        "best_source": result_source,
        "best_source_id": result_source_id,
        "candidate_document_html": (
            best_evidence["candidate_document_html"]
            if best_evidence is not None
            else (
                passage_groups[0]["candidate_document_html"]
                if passage_groups
                else html.escape(candidate)
            )
        ),
        "tfidf_score": result_tfidf,
        "semantic_score": result_semantic,
        "hybrid_score": result_hybrid,
        "novelty_score": round(max(0.0, 1 - result_hybrid), 4),
        "contextual_analysis": contextual_for_response,
        "contextual_explanation": explanation_for_response,
        "sources": displayed_sources[:TOP_SOURCES],
        "matches": passage_groups,
    }


@asynccontextmanager
async def lifespan(app: FastAPI):
    """Initialise le modèle et les index au démarrage de l'API."""
    model_info = get_model_info()
    print(f"[similarity] Modèle : {model_info['name']}")

    get_model()

    def initialize_reference_data():
        # Django ORM et call_command('migrate') sont synchrones.
        # L'initialisation est donc exécutée dans un thread.
        stats = ensure_reference_index()
        _warm_reference_runtime_indexes()
        return stats

    try:
        stats = await run_in_threadpool(initialize_reference_data)

        print(
            "[similarity] Références : "
            f"{stats['created_or_updated']} préparée(s), "
            f"{stats['skipped']} déjà indexée(s)."
        )

    except Exception as exc:
        print(
            "[similarity] Initialisation PostgreSQL/index impossible : "
            f"{exc}"
        )

    yield


app = FastAPI(
    title="TDRDOC-SCAN - Détection documentaire",
    version="2.0.0",
    lifespan=lifespan,
)


@app.get("/")
def root():
    return {
        "service": "tdrdoc-scan",
        "status": "ok",
        "version": "2.0.0",
    }


@app.get("/api/health")
def health():
    result = {
        "status": "ok",
        "model": get_model_info(),
        "reference_counts": {},
        "cache_ready": list(_REFERENCE_CACHE.keys()),
    }

    for kind in ("duplicate", "plagiarism"):
        try:
            result["reference_counts"][kind] = len(
                load_reference_documents(kind)[0]
            )
        except Exception:
            result["reference_counts"][kind] = 0

    return result


@app.post("/api/detect/{kind}")
async def detect(
    kind: str,
    file: UploadFile = File(...),
    threshold: float | None = Form(None),
):
    if kind not in ("plagiarism", "duplicate"):
        raise HTTPException(
            status_code=404,
            detail="Mode inconnu. Utilisez plagiarism ou duplicate.",
        )

    if (
        not file.filename
        or Path(file.filename).suffix.lower() not in {".docx", ".pdf"}
    ):
        raise HTTPException(
            status_code=400,
            detail="Formats acceptés : DOCX et PDF.",
        )

    start = time.perf_counter()
    safe_name = re.sub(r"[^\w.\- ]", "_", Path(file.filename).name)

    with tempfile.TemporaryDirectory() as directory:
        path = Path(directory) / safe_name
        path.write_bytes(await file.read())

        candidate_text = read_document(path)

        # run_detection utilise le Django ORM de façon synchrone.
        result = await run_in_threadpool(
            run_detection,
            candidate_text,
            kind,
            threshold,
        )

    result["duration_ms"] = round(
        (time.perf_counter() - start) * 1000
    )

    return result


@app.post("/api/compare")
async def compare(
    file_a: UploadFile = File(...),
    file_b: UploadFile = File(...),
):
    for uploaded_file in (file_a, file_b):
        if (
            not uploaded_file.filename
            or Path(uploaded_file.filename).suffix.lower()
            not in {".docx", ".pdf"}
        ):
            raise HTTPException(
                status_code=400,
                detail=(
                    "Les deux fichiers doivent être au format DOCX ou PDF."
                ),
            )

    with tempfile.TemporaryDirectory() as directory:
        suffix_a = Path(file_a.filename).suffix.lower()
        suffix_b = Path(file_b.filename).suffix.lower()

        path_a = Path(directory) / f"a{suffix_a}"
        path_b = Path(directory) / f"b{suffix_b}"

        path_a.write_bytes(await file_a.read())
        path_b.write_bytes(await file_b.read())

        text_a = read_document(path_a)
        text_b = read_document(path_b)

        result = score_two_texts(text_a, text_b)

    if result is None:
        raise HTTPException(
            status_code=400,
            detail="Les deux documents doivent contenir du texte exploitable.",
        )

    return result


@app.post("/api/tdr-rag/analyze")
async def analyze_tdr_rag_independent(file: UploadFile = File(...)):
    """Point d'entrée dédié au moteur RAG indépendant des TDR."""
    if (
        not file.filename
        or Path(file.filename).suffix.lower() not in {".docx", ".pdf"}
    ):
        raise HTTPException(
            status_code=400,
            detail="Formats acceptés : DOCX et PDF.",
        )

    from src.tdr_rag.independent_engine import (
        analyze_tdr_independently,
        extract_tdr_document,
        IndependentTDRError,
    )

    try:
        with tempfile.TemporaryDirectory() as directory:
            safe_name = re.sub(
                r"[^\w.\- ]",
                "_",
                Path(file.filename).name,
            )

            path = Path(directory) / safe_name
            path.write_bytes(await file.read())

            candidate_text = extract_tdr_document(path)

        result = await run_in_threadpool(
            analyze_tdr_independently,
            candidate_text,
        )

        return result

    except IndependentTDRError as exc:
        raise HTTPException(
            status_code=503,
            detail=str(exc),
        ) from exc
