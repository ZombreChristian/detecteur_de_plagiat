
"""
Extraction du texte des fichiers DOCX du corpus.

Les fichiers DOCX possèdent des noms longs et parfois tronqués.
Le projet utilise cependant des identifiants normalisés :

    TDR_001_....docx      -> TDR-001
    Rapport_001_....docx  -> RAP-001

Entrée :
    donnees/TDR/*.docx
    donnees/Rapport d'etude/*.docx

Sortie :
    donnees/documents_extraits.csv
"""

from pathlib import Path
import re

from docx import Document
import pandas as pd


# ============================================================
# CHEMINS
# ============================================================

ROOT = Path(__file__).resolve().parents[2]

DATA = ROOT / "donnees"

OUT = DATA / "documents_extraits.csv"


# ============================================================
# EXTRACTION DU TEXTE DOCX
# ============================================================

def extract_docx(path: Path) -> str:
    """
    Extrait le texte des paragraphes et des tableaux d'un DOCX.
    """

    doc = Document(path)

    parts = []

    # Paragraphes
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()

        if text:
            parts.append(text)

    # Tableaux
    for table in doc.tables:
        for row in table.rows:

            cells = [
                cell.text.strip()
                for cell in row.cells
            ]

            if any(cells):
                parts.append(" | ".join(cells))

    return "\n".join(parts)


# ============================================================
# IDENTIFIANT NORMALISÉ
# ============================================================

def build_document_id(path: Path, document_type: str) -> str:
    """
    Construit l'identifiant canonique du document.

    Exemples :

        TDR_001_Diagnostic....docx
            -> TDR-001

        Rapport_001_Diagnostic....docx
            -> RAP-001
    """

    # Recherche du premier numéro à trois chiffres
    match = re.search(r"(\d{3})", path.stem)

    if not match:
        raise ValueError(
            f"Impossible de déterminer le numéro du document : "
            f"{path.name}"
        )

    number = match.group(1)

    if document_type == "TDR":
        return f"TDR-{number}"

    if document_type == "Rapport":
        return f"RAP-{number}"

    raise ValueError(
        f"Type de document inconnu : {document_type}"
    )


# ============================================================
# PROGRAMME PRINCIPAL
# ============================================================

def main():

    print("=" * 70)
    print("EXTRACTION DU CORPUS DOCX")
    print("=" * 70)

    DATA.mkdir(
        parents=True,
        exist_ok=True
    )

    rows = []

    # ========================================================
    # TDR
    # ========================================================

    tdr_folder = DATA / "TDR"

    tdr_files = sorted(
        tdr_folder.glob("*.docx")
    )

    print(f"\nTDR trouvés : {len(tdr_files)}")

    for path in tdr_files:

        document_id = build_document_id(
            path,
            "TDR"
        )

        text = extract_docx(path)

        rows.append(
            {
                "document_id": document_id,
                "type_document": "TDR",
                "path": str(
                    path.relative_to(ROOT)
                ),
                "text": text,
                "nb_caracteres": len(text),
                "nb_mots": len(text.split())
            }
        )

    # ========================================================
    # RAPPORTS
    # ========================================================

    reports_folder = DATA / "Rapport d'etude"

    report_files = sorted(
        reports_folder.glob("*.docx")
    )

    print(f"Rapports trouvés : {len(report_files)}")

    for path in report_files:

        document_id = build_document_id(
            path,
            "Rapport"
        )

        text = extract_docx(path)

        rows.append(
            {
                "document_id": document_id,
                "type_document": "Rapport",
                "path": str(
                    path.relative_to(ROOT)
                ),
                "text": text,
                "nb_caracteres": len(text),
                "nb_mots": len(text.split())
            }
        )

    # ========================================================
    # DATAFRAME
    # ========================================================

    df = pd.DataFrame(rows)

    # ========================================================
    # VÉRIFICATIONS
    # ========================================================

    print("\nVérification des identifiants...")

    # Vérification des doublons
    duplicates = df[
        df["document_id"].duplicated(
            keep=False
        )
    ]

    if not duplicates.empty:

        print(
            "\nERREUR : identifiants dupliqués :"
        )

        print(
            duplicates[
                ["document_id", "path"]
            ].to_string(index=False)
        )

        raise ValueError(
            "Des document_id sont dupliqués."
        )

    # Vérification des quantités
    nb_tdr = (
        df["document_id"]
        .str.startswith("TDR-")
        .sum()
    )

    nb_rap = (
        df["document_id"]
        .str.startswith("RAP-")
        .sum()
    )

    print(f"Documents TDR   : {nb_tdr}")
    print(f"Documents RAP   : {nb_rap}")
    print(f"Total documents : {len(df)}")

    # Le corpus attendu contient 60 TDR + 60 rapports
    if nb_tdr != 60:
        raise ValueError(
            f"Nombre de TDR inattendu : {nb_tdr}. "
            f"Le corpus doit contenir 60 TDR."
        )

    if nb_rap != 60:
        raise ValueError(
            f"Nombre de rapports inattendu : {nb_rap}. "
            f"Le corpus doit contenir 60 rapports."
        )

    # ========================================================
    # TRI
    # ========================================================

    df["numero"] = (
        df["document_id"]
        .str.extract(r"(\d+)")
        .astype(int)
    )

    df["ordre_type"] = df["document_id"].apply(
        lambda x: 0 if x.startswith("TDR-") else 1
    )

    df = (
        df
        .sort_values(
            ["ordre_type", "numero"]
        )
        .drop(
            columns=["numero", "ordre_type"]
        )
        .reset_index(drop=True)
    )

    # ========================================================
    # SAUVEGARDE
    # ========================================================

    df.to_csv(
        OUT,
        index=False,
        encoding="utf-8-sig"
    )

    # ========================================================
    # RÉSUMÉ
    # ========================================================

    print("\nFichier créé :")
    print(OUT)

    print("\nAperçu des premiers documents :")

    print(
        df[
            [
                "document_id",
                "type_document",
                "path",
                "nb_mots"
            ]
        ]
        .head(10)
        .to_string(index=False)
    )

    print("\nAperçu des rapports :")

    print(
        df[
            [
                "document_id",
                "type_document",
                "path",
                "nb_mots"
            ]
        ]
        .tail(10)
        .to_string(index=False)
    )

    print("\n" + "=" * 70)
    print("EXTRACTION TERMINÉE AVEC SUCCÈS")
    print("=" * 70)


# ============================================================
# POINT D'ENTRÉE
# ============================================================

if __name__ == "__main__":
    main()
