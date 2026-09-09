
"""
Prédiction de similarité entre deux documents DOCX.

Le programme calcule :
    - similarité TF-IDF
    - similarité sémantique
    - score hybride

Puis utilise le modèle de classification entraîné
pour produire la probabilité que les deux documents
soient similaires.

Usage :

python -m src.prediction.predict \
    donnees/TDR/TDR_001.docx \
    donnees/TDR/TDR_002.docx
"""

from pathlib import Path

import joblib
import numpy as np

from docx import Document

from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

from sentence_transformers import SentenceTransformer


# ============================================================
# CONFIGURATION
# ============================================================

ROOT = Path(__file__).resolve().parents[2]

MODEL_PATH = (
    ROOT
    / "modeles"
    / "similarity_classifier.joblib"
)

MODEL_NAME = (
    "sentence-transformers/"
    "paraphrase-multilingual-MiniLM-L12-v2"
)


# ============================================================
# LECTURE DU DOCUMENT DOCX
# ============================================================

def read_docx(path: Path) -> str:
    """
    Extrait le texte des paragraphes et des tableaux d'un DOCX.
    """

    if not path.exists():
        raise FileNotFoundError(
            f"Document introuvable : {path}"
        )

    if path.suffix.lower() != ".docx":
        raise ValueError(
            f"Le fichier doit être un DOCX : {path}"
        )

    document = Document(path)

    parts = []

    # Paragraphes
    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            parts.append(text)

    # Tableaux
    for table in document.tables:

        for row in table.rows:

            cells = []

            for cell in row.cells:

                text = cell.text.strip()

                if text:
                    cells.append(text)

            if cells:
                parts.append(" | ".join(cells))

    return " ".join(parts)


# ============================================================
# CALCUL TF-IDF
# ============================================================

def calculate_tfidf_similarity(text_a: str, text_b: str) -> float:

    vectorizer = TfidfVectorizer(
        ngram_range=(1, 2),
        sublinear_tf=True
    )

    matrix = vectorizer.fit_transform(
        [text_a, text_b]
    )

    score = cosine_similarity(
        matrix[0],
        matrix[1]
    )[0, 0]

    return float(score)


# ============================================================
# CALCUL SIMILARITÉ SÉMANTIQUE
# ============================================================

def calculate_semantic_similarity(
    text_a: str,
    text_b: str
) -> float:

    model = SentenceTransformer(
        MODEL_NAME
    )

    embeddings = model.encode(
        [text_a, text_b],
        normalize_embeddings=True,
        show_progress_bar=False
    )

    score = np.dot(
        embeddings[0],
        embeddings[1]
    )

    return float(score)


# ============================================================
# PRÉDICTION
# ============================================================

def predict(file_a, file_b):

    file_a = Path(file_a)
    file_b = Path(file_b)

    print("\nLecture des documents...")

    text_a = read_docx(file_a)
    text_b = read_docx(file_b)

    if not text_a:
        raise ValueError(
            f"Le document {file_a.name} ne contient aucun texte exploitable."
        )

    if not text_b:
        raise ValueError(
            f"Le document {file_b.name} ne contient aucun texte exploitable."
        )

    # --------------------------------------------------------
    # TF-IDF
    # --------------------------------------------------------

    print("Calcul de la similarité TF-IDF...")

    tfidf_score = calculate_tfidf_similarity(
        text_a,
        text_b
    )

    # --------------------------------------------------------
    # Embeddings
    # --------------------------------------------------------

    print("Calcul de la similarité sémantique...")

    semantic_score = calculate_semantic_similarity(
        text_a,
        text_b
    )

    # --------------------------------------------------------
    # Score hybride
    # --------------------------------------------------------

    hybrid_score = (
        0.35 * tfidf_score
        + 0.65 * semantic_score
    )

    # --------------------------------------------------------
    # Chargement du classifieur
    # --------------------------------------------------------

    if not MODEL_PATH.exists():

        raise FileNotFoundError(
            f"Modèle introuvable : {MODEL_PATH}\n\n"
            "Tu dois d'abord exécuter :\n"
            "python src/training/train_model.py"
        )

    classifier = joblib.load(
        MODEL_PATH
    )

    # --------------------------------------------------------
    # Construction des caractéristiques
    # --------------------------------------------------------

    X = np.array([
        [
            tfidf_score,
            semantic_score,
            hybrid_score
        ]
    ])

    # --------------------------------------------------------
    # Prédiction ML
    # --------------------------------------------------------

    prediction = int(
        classifier.predict(X)[0]
    )

    probability = float(
        classifier.predict_proba(X)[0][1]
    )

    # --------------------------------------------------------
    # Décision
    # --------------------------------------------------------

    if probability >= 0.70:

        decision = "SIMILAIRES"

    elif probability >= 0.50:

        decision = "A EXAMINER"

    else:

        decision = "DIFFERENTS"

    # --------------------------------------------------------
    # Résultat
    # --------------------------------------------------------

    result = {

        "document_a": str(file_a),

        "document_b": str(file_b),

        "tfidf_score": round(
            tfidf_score,
            4
        ),

        "semantic_score": round(
            semantic_score,
            4
        ),

        "hybrid_score": round(
            hybrid_score,
            4
        ),

        "similarity_probability": round(
            probability,
            4
        ),

        "prediction": prediction,

        "decision": decision
    }

    return result


# ============================================================
# AFFICHAGE
# ============================================================

def display_result(result):

    print("\n")
    print("=" * 70)
    print("RÉSULTAT DE LA DÉTECTION")
    print("=" * 70)

    print(
        f"\nDocument A : {result['document_a']}"
    )

    print(
        f"Document B : {result['document_b']}"
    )

    print("\nScores :")

    print(
        f"  TF-IDF             : "
        f"{result['tfidf_score']:.4f}"
    )

    print(
        f"  Similarité sémantique : "
        f"{result['semantic_score']:.4f}"
    )

    print(
        f"  Score hybride      : "
        f"{result['hybrid_score']:.4f}"
    )

    print(
        f"\nProbabilité de similarité : "
        f"{result['similarity_probability']:.2%}"
    )

    print(
        f"\nPrédiction ML : "
        f"{result['prediction']}"
    )

    print(
        f"Décision : "
        f"{result['decision']}"
    )

    print("\n" + "=" * 70)


# ============================================================
# PROGRAMME PRINCIPAL
# ============================================================

if __name__ == "__main__":

    import argparse

    parser = argparse.ArgumentParser(
        description=(
            "Détection de similarité entre deux "
            "documents DOCX."
        )
    )

    parser.add_argument(
        "document_a",
        help="Premier document DOCX"
    )

    parser.add_argument(
        "document_b",
        help="Deuxième document DOCX"
    )

    args = parser.parse_args()

    result = predict(
        args.document_a,
        args.document_b
    )

    display_result(result)

